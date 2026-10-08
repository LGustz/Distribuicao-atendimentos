# ============================================================
# CENTRAL DE VISTORIAS
# Distribuição automática inteligente
#
# REGRA PRINCIPAL:
#   Cada técnico pode receber até 2 vistorias em TODOS
#   os horários.
#
# Requisito:
#   python -m pip install openpyxl
#
# Uso:
#   python distribuir.py Entrada.xlsx saida.xlsx
#
# Entrada:
#   .xlsx
#   .xlsm
#   .csv
#   .txt
#
# Colunas esperadas:
#   1 = SINISTRO
#   2 = HORÁRIO
#   3 = ASSISTÊNCIA_COD
# ============================================================

import sys
import csv
import re
from pathlib import Path
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURAÇÃO
# ============================================================

CONFIG = {
    # Bloqueia sinistros repetidos.
    "bloquear_duplicados": True,

    # Considera somente técnicos com active=True.
    "respeitar_active": True,

    # Mantém técnicos com menor janela disponíveis quando
    # possível.
    "preservar_janelas_restritas": True,

    # Pesos utilizados no balanceamento.
    "peso_carga": 1000,
    "peso_utilizacao": 500,
    "peso_escassez": 300,
    "peso_horario": 50,
}


# ============================================================
# TÉCNICOS
# ============================================================

TECHS = [
    {
        "name": "João Pedro",
        "type": "Integral",
        "start": "08:30",
        "end": "17:50",
        "active": True,
    },
    {
        "name": "Kevin",
        "type": "Integral",
        "start": "08:30",
        "end": "17:50",
        "active": True,
    },
    {
        "name": "Guilherme",
        "type": "Integral",
        "start": "08:30",
        "end": "17:50",
        "active": True,
    },
    {
        "name": "Nicolas01",
        "type": "Integral",
        "start": "08:30",
        "end": "18:50",
        "active": True,
    },
    {
        "name": "Luiz",
        "type": "Integral",
        "start": "10:30",
        "end": "17:50",
        "active": True,
    },
    {
        "name": "Isaque",
        "type": "Integral",
        "start": "08:30",
        "end": "17:30",
        "active": True,
    },
    {
        "name": "Mateus Gabriel",
        "type": "Estagiário",
        "start": "08:30",
        "end": "13:50",
        "active": True,
    },
    {
        "name": "Matheus Salgado",
        "type": "Estagiário",
        "start": "08:30",
        "end": "13:50",
        "active": True,
    },
    {
        "name": "Caua Gabriel",
        "type": "Estagiário",
        "start": "14:10",
        "end": "18:50",
        "active": True,
    },
    {
        "name": "Matheu Henrique",
        "type": "Estagiário",
        "start": "14:10",
        "end": "18:50",
        "active": True,
    },
    {
        "name": "Nicolas02",
        "type": "Estagiário",
        "start": "14:10",
        "end": "18:50",
        "active": True,
    },
    {
        "name": "Ryan",
        "type": "Estagiário",
        "start": "14:10",
        "end": "18:50",
        "active": True,
    },
]


# ============================================================
# HORÁRIOS
# ============================================================

def to_min(time_str):
    """
    Converte HH:MM para minutos.
    Exemplo:
        08:30 -> 510
    """

    hour, minute = time_str.split(":")[:2]

    return int(hour) * 60 + int(minute)


def from_min(minutes):
    """
    Converte minutos para HH:MM.
    """

    return (
        f"{minutes // 60:02d}:"
        f"{minutes % 60:02d}"
    )


def generate_times(start, end, step=20):
    """
    Gera horários de start até end.
    """

    current = to_min(start)
    final = to_min(end)

    result = []

    while current <= final:

        result.append(
            from_min(current)
        )

        current += step

    return result


# ============================================================
# GRADE DE HORÁRIOS
# ============================================================

# Manhã:
# 08:30 até 11:30
#
# Tarde:
# 13:10 até 18:30
#
# Último horário:
# 18:50
#
# TODOS possuem limite de 2 vistorias por técnico.

TIMES = (
    generate_times(
        "08:30",
        "11:30",
        20,
    )
    +
    generate_times(
        "13:10",
        "18:30",
        20,
    )
    +
    ["18:50"]
)

# Remove qualquer eventual duplicidade.
TIMES = list(
    dict.fromkeys(TIMES)
)


# ============================================================
# LIMITE DE VISTORIAS
# ============================================================

def slot_limit(time):
    """
    REGRA OPERACIONAL:

    Todo técnico pode receber no máximo
    2 vistorias em qualquer horário.

    Não existem mais horários especiais.
    """

    return 2


# ============================================================
# NORMALIZAÇÃO
# ============================================================

def normalize_id(value):
    """
    Normaliza o número do sinistro.
    """

    if value is None:
        return ""

    if isinstance(value, float):

        if value.is_integer():
            return str(int(value))

    text = str(value).strip()

    if text.endswith(".0"):
        text = text[:-2]

    return text


def norm_time(value):
    """
    Normaliza horários vindos do Excel,
    CSV ou TXT.

    Aceita:
        08:30
        8:30
        08h30
        08.30
        08:30:00
        datetime.time
        datetime.datetime
        fração de dia do Excel
    """

    if value is None:
        return ""

    # Excel datetime/time
    if hasattr(value, "hour") and hasattr(
        value,
        "minute",
    ):

        return (
            f"{value.hour:02d}:"
            f"{value.minute:02d}"
        )

    # Excel pode armazenar horário como
    # fração do dia.
    if isinstance(value, (int, float)):

        if 0 <= value < 1:

            minutes = round(
                value * 24 * 60
            )

            return from_min(
                minutes
            )

    text = str(value).strip()

    match = re.match(
        r"^(\d{1,2})[:h.](\d{2})(?::\d{2})?$",
        text,
        re.IGNORECASE,
    )

    if match:

        hour = int(
            match.group(1)
        )

        minute = int(
            match.group(2)
        )

        if (
            0 <= hour <= 23
            and
            0 <= minute <= 59
        ):

            return (
                f"{hour:02d}:"
                f"{minute:02d}"
            )

    return text


# ============================================================
# LEITURA DE ARQUIVOS
# ============================================================

def read_excel(path):

    workbook = load_workbook(
        path,
        data_only=True,
        read_only=True,
    )

    worksheet = workbook.active

    rows = [
        list(row)
        for row in worksheet.iter_rows(
            values_only=True
        )
    ]

    workbook.close()

    return rows


def detect_delimiter(text):

    try:

        dialect = csv.Sniffer().sniff(
            text,
            delimiters=";,\t",
        )

        return dialect.delimiter

    except csv.Error:

        return ";"


def read_text_file(path):

    with open(
        path,
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        text = file.read()

    if not text.strip():
        return []

    delimiter = detect_delimiter(
        text[:5000]
    )

    reader = csv.reader(
        text.splitlines(),
        delimiter=delimiter,
    )

    return list(reader)


def read_rows(path):

    extension = (
        Path(path)
        .suffix
        .lower()
    )

    if extension in {
        ".xlsx",
        ".xlsm",
    }:

        return read_excel(path)

    if extension in {
        ".csv",
        ".txt",
    }:

        return read_text_file(path)

    raise ValueError(
        f"Formato não suportado: {extension}. "
        "Use .xlsx, .xlsm, .csv ou .txt."
    )


# ============================================================
# IMPORTAÇÃO DOS SINISTROS
# ============================================================

def parse_input(path):

    items = []
    invalid = []
    duplicates = []

    seen_ids = set()

    rows = read_rows(path)

    for line_number, row in enumerate(
        rows,
        start=1,
    ):

        cells = list(row) + [
            None,
            None,
            None,
        ]

        sinistro = normalize_id(
            cells[0]
        )

        raw_time = cells[1]

        assist = (
            str(cells[2]).strip()
            if cells[2] not in (
                None,
                "",
            )
            else "—"
        )

        # Ignora linha totalmente vazia.
        if (
            not sinistro
            and not raw_time
            and not cells[2]
        ):
            continue

        # Detecta cabeçalho.
        header_text = " ".join(
            str(x or "")
            for x in cells[:3]
        ).upper()

        if "SINISTRO" in header_text:
            continue

        # Sinistro vazio.
        if not sinistro:

            invalid.append(
                {
                    "line": line_number,
                    "reason": "Sinistro vazio",
                    "id": "",
                    "time": "",
                    "assist": assist,
                }
            )

            continue

        # Sinistro duplicado.
        if (
            CONFIG["bloquear_duplicados"]
            and
            sinistro in seen_ids
        ):

            duplicates.append(
                {
                    "line": line_number,
                    "reason": "Sinistro duplicado",
                    "id": sinistro,
                    "time": norm_time(
                        raw_time
                    ),
                    "assist": assist,
                }
            )

            continue

        seen_ids.add(
            sinistro
        )

        time = norm_time(
            raw_time
        )

        # Horário inválido.
        if time not in TIMES:

            invalid.append(
                {
                    "line": line_number,
                    "reason": (
                        f'Horário inválido: '
                        f'"{raw_time}"'
                    ),
                    "id": sinistro,
                    "time": time,
                    "assist": assist,
                }
            )

            continue

        items.append(
            {
                "id": sinistro,
                "time": time,
                "assist": assist,
                "line": line_number,
            }
        )

    return (
        items,
        invalid,
        duplicates,
    )


# ============================================================
# TÉCNICOS
# ============================================================

def get_active_techs():

    if not CONFIG[
        "respeitar_active"
    ]:

        return list(TECHS)

    return [
        tech
        for tech in TECHS
        if tech.get(
            "active",
            True,
        )
    ]


def validate_techs(techs):

    errors = []

    names = set()

    for tech in techs:

        name = tech.get(
            "name",
            "",
        ).strip()

        if not name:

            errors.append(
                "Existe técnico sem nome."
            )

            continue

        if name in names:

            errors.append(
                f"Técnico duplicado: {name}"
            )

        names.add(name)

        try:

            start = to_min(
                tech["start"]
            )

            end = to_min(
                tech["end"]
            )

            if start > end:

                errors.append(
                    f"Horário inválido de "
                    f"{name}: "
                    f"{tech['start']} > "
                    f"{tech['end']}"
                )

        except Exception:

            errors.append(
                f"Horário inválido do "
                f"técnico {name}."
            )

        if tech.get("type") not in {
            "Integral",
            "Estagiário",
        }:

            errors.append(
                f"Tipo inválido de {name}: "
                f"{tech.get('type')}"
            )

    return errors


# ============================================================
# ELEGIBILIDADE
# ============================================================

def eligible(
    tech,
    time,
):
    """
    Verifica se o técnico pode receber
    uma vistoria naquele horário.
    """

    minutes = to_min(
        time
    )

    start = to_min(
        tech["start"]
    )

    end = to_min(
        tech["end"]
    )

    if minutes < start:
        return False

    if minutes > end:
        return False

    # Somente integrais no horário 16:10.
    #
    # Essa regra continua existindo.
    # Se quiser remover depois, basta alterar
    # INTEGRAL_ONLY.
    if (
        time in INTEGRAL_ONLY
        and
        tech["type"] != "Integral"
    ):

        return False

    return True


# ============================================================
# REGRAS ESPECIAIS DE HORÁRIO
# ============================================================

# Horário em que somente técnicos integrais
# podem receber vistorias.
#
# Se NÃO quiser essa restrição, deixe:
#
# INTEGRAL_ONLY = set()

INTEGRAL_ONLY = {
    "16:10",
}


# ============================================================
# CAPACIDADE
# ============================================================

def calculate_capacity(
    tech
):
    """
    Calcula a capacidade total do técnico.

    Cada horário elegível = 2 vistorias.
    """

    total = 0
    slots = 0

    for time in TIMES:

        if eligible(
            tech,
            time,
        ):

            total += slot_limit(
                time
            )

            slots += 1

    return (
        total,
        slots,
    )


def calculate_future_capacity(
    tech,
    current_index,
):
    """
    Calcula quantas vistorias o técnico
    ainda poderia receber nos próximos horários.
    """

    total = 0

    for time in TIMES[
        current_index + 1:
    ]:

        if eligible(
            tech,
            time,
        ):

            total += slot_limit(
                time
            )

    return total


def calculate_available_techs_by_time(
    techs
):

    result = {}

    for time in TIMES:

        result[time] = [
            tech
            for tech in techs
            if eligible(
                tech,
                time,
            )
        ]

    return result


# ============================================================
# DISTRIBUIÇÃO
# ============================================================

def distribute(
    items,
    techs,
):
    """
    Distribui os sinistros procurando:

    1. Respeitar horário do técnico.
    2. Respeitar regras de tipo.
    3. Máximo de 2 vistorias por técnico
       em QUALQUER horário.
    4. Balancear carga entre técnicos.
    5. Preservar técnicos com menor janela.
    6. Priorizar horários com maior pressão.
    """

    assignments = {
        tech["name"]: {
            time: []
            for time in TIMES
        }
        for tech in techs
    }

    load = {
        tech["name"]: 0
        for tech in techs
    }

    capacity = {}

    available_techs = (
        calculate_available_techs_by_time(
            techs
        )
    )

    for tech in techs:

        cap, _ = calculate_capacity(
            tech
        )

        capacity[
            tech["name"]
        ] = cap

    # --------------------------------------------------------
    # Agrupa sinistros por horário.
    # --------------------------------------------------------

    by_time = defaultdict(list)

    for item in items:

        by_time[
            item["time"]
        ].append(item)

    demand = {
        time: len(
            by_time[time]
        )
        for time in TIMES
    }

    # --------------------------------------------------------
    # Calcula pressão de cada horário.
    # --------------------------------------------------------

    time_priority = {}

    for index, time in enumerate(
        TIMES
    ):

        available_count = len(
            available_techs[time]
        )

        capacity_at_time = (
            available_count
            *
            slot_limit(time)
        )

        demand_at_time = demand[
            time
        ]

        if capacity_at_time:

            pressure = (
                demand_at_time
                /
                capacity_at_time
            )

        else:

            pressure = (
                float("inf")
                if demand_at_time
                else 0
            )

        time_priority[time] = (
            -pressure,
            available_count,
            index,
        )

    ordered_times = sorted(
        TIMES,
        key=lambda x:
            time_priority[x],
    )

    # --------------------------------------------------------
    # Distribuição.
    # --------------------------------------------------------

    assignment_log = []

    for time in ordered_times:

        items_at_time = by_time[
            time
        ]

        if not items_at_time:
            continue

        current_index = TIMES.index(
            time
        )

        for item in items_at_time:

            candidates = []

            for tech in available_techs[
                time
            ]:

                name = tech[
                    "name"
                ]

                # ------------------------------------------------
                # LIMITE:
                # 2 vistorias em TODOS os horários.
                # ------------------------------------------------

                if len(
                    assignments[name][time]
                ) >= 2:

                    continue

                total_capacity = capacity[
                    name
                ]

                used = load[
                    name
                ]

                utilization = (
                    used
                    /
                    total_capacity
                    if total_capacity
                    else 1
                )

                future_capacity = (
                    calculate_future_capacity(
                        tech,
                        current_index,
                    )
                )

                # ------------------------------------------------
                # Carga futura.
                # ------------------------------------------------

                future_demand = 0

                for future_time in TIMES[
                    current_index + 1:
                ]:

                    future_demand += demand[
                        future_time
                    ]

                # ------------------------------------------------
                # Escassez.
                # ------------------------------------------------

                if CONFIG[
                    "preservar_janelas_restritas"
                ]:

                    scarcity = (
                        1
                        /
                        max(
                            future_capacity,
                            1,
                        )
                    )

                else:

                    scarcity = 0

                # ------------------------------------------------
                # Pontuação.
                # ------------------------------------------------

                load_score = (
                    utilization
                    *
                    CONFIG[
                        "peso_utilizacao"
                    ]
                )

                absolute_load_score = (
                    used
                    *
                    CONFIG[
                        "peso_carga"
                    ]
                    /
                    100
                )

                scarcity_score = (
                    scarcity
                    *
                    CONFIG[
                        "peso_escassez"
                    ]
                )

                if future_capacity:

                    future_pressure = (
                        future_demand
                        /
                        future_capacity
                    )

                else:

                    future_pressure = 0

                future_pressure_score = (
                    future_pressure
                    *
                    CONFIG[
                        "peso_horario"
                    ]
                )

                score = (
                    load_score
                    +
                    absolute_load_score
                    +
                    scarcity_score
                    +
                    future_pressure_score
                )

                # Pequeno desempate em favor de
                # técnicos com maior capacidade.
                score -= (
                    total_capacity
                    * 0.001
                )

                candidates.append(
                    {
                        "score": score,
                        "name": name,
                        "tech": tech,
                        "future_capacity":
                            future_capacity,
                        "utilization":
                            utilization,
                    }
                )

            # ----------------------------------------------------
            # Nenhum candidato.
            # ----------------------------------------------------

            if not candidates:

                assignment_log.append(
                    {
                        "item": item,
                        "time": time,
                        "status":
                            "NÃO DISTRIBUÍDO",
                        "reason":
                            "Capacidade do horário "
                            "esgotada",
                        "tech": "",
                    }
                )

                continue

            # ----------------------------------------------------
            # Escolhe menor score.
            # ----------------------------------------------------

            candidates.sort(
                key=lambda candidate: (
                    candidate["score"],
                    candidate["name"],
                )
            )

            selected = candidates[0]

            name = selected[
                "name"
            ]

            assignments[
                name
            ][time].append(
                item
            )

            load[
                name
            ] += 1

            assignment_log.append(
                {
                    "item": item,
                    "time": time,
                    "status":
                        "DISTRIBUÍDO",
                    "reason": "OK",
                    "tech": name,
                    "score":
                        selected["score"],
                }
            )

    # --------------------------------------------------------
    # Descobre não distribuídos.
    # --------------------------------------------------------

    assigned_ids = {
        entry["item"]["id"]
        for entry in assignment_log
        if entry["status"]
        == "DISTRIBUÍDO"
    }

    unassigned = []

    for item in items:

        if item["id"] not in assigned_ids:

            has_eligible = any(
                eligible(
                    tech,
                    item["time"],
                )
                for tech in techs
            )

            if has_eligible:

                reason = (
                    "Capacidade do horário "
                    "esgotada"
                )

            else:

                reason = (
                    "Nenhum técnico elegível "
                    "para este horário"
                )

            unassigned.append(
                {
                    **item,
                    "reason": reason,
                }
            )

    return (
        assignments,
        load,
        capacity,
        demand,
        available_techs,
        unassigned,
        assignment_log,
    )


# ============================================================
# ESTATÍSTICAS
# ============================================================

def calculate_stats(
    techs,
    load,
    capacity,
    demand,
    available_techs,
    unassigned,
):

    technician_stats = {}

    for tech in techs:

        name = tech[
            "name"
        ]

        cap = capacity[
            name
        ]

        used = load[
            name
        ]

        technician_stats[
            name
        ] = {
            "capacity": cap,
            "used": used,
            "remaining":
                max(
                    cap - used,
                    0,
                ),
            "utilization":
                (
                    used
                    /
                    cap
                    *
                    100
                    if cap
                    else 0
                ),
        }

    time_stats = {}

    for time in TIMES:

        available_count = len(
            available_techs[
                time
            ]
        )

        capacity_at_time = (
            available_count
            *
            2
        )

        used = demand[
            time
        ]

        remaining = (
            capacity_at_time
            -
            used
        )

        utilization = (
            used
            /
            capacity_at_time
            *
            100
            if capacity_at_time
            else 0
        )

        if used == 0:

            status = (
                "SEM DEMANDA"
            )

        elif remaining < 0:

            status = (
                "SOBRECARGA"
            )

        elif remaining == 0:

            status = (
                "NO LIMITE"
            )

        elif remaining <= 2:

            status = (
                "CRÍTICO"
            )

        else:

            status = "OK"

        time_stats[
            time
        ] = {
            "demand": used,
            "available":
                available_count,
            "capacity":
                capacity_at_time,
            "remaining":
                remaining,
            "utilization":
                utilization,
            "status":
                status,
        }

    total = sum(
        demand.values()
    )

    distributed = (
        total
        -
        len(unassigned)
    )

    total_capacity = sum(
        capacity.values()
    )

    return {
        "technicians":
            technician_stats,
        "times":
            time_stats,
        "total":
            total,
        "distributed":
            distributed,
        "unassigned":
            len(unassigned),
        "capacity":
            total_capacity,
        "distribution_rate":
            (
                distributed
                /
                total
                *
                100
                if total
                else 0
            ),
    }


# ============================================================
# ESTILOS EXCEL
# ============================================================

HEAD = PatternFill(
    "solid",
    start_color="1E293B",
)

ALT = PatternFill(
    "solid",
    start_color="F8FAFC",
)

EXIT_FILL = PatternFill(
    "solid",
    start_color="FEF3C7",
)

GREEN = PatternFill(
    "solid",
    start_color="DCFCE7",
)

YELLOW = PatternFill(
    "solid",
    start_color="FEF9C3",
)

RED = PatternFill(
    "solid",
    start_color="FEE2E2",
)

BLUE = PatternFill(
    "solid",
    start_color="DBEAFE",
)

WHITE = Font(
    name="Arial",
    bold=True,
    color="FFFFFF",
)

BASE = Font(
    name="Arial",
    size=10,
)

BOLD = Font(
    name="Arial",
    bold=True,
)

THIN = Side(
    style="thin",
    color="CBD5E1",
)

BOX = Border(
    left=THIN,
    right=THIN,
    top=THIN,
    bottom=THIN,
)


def style_header(
    ws,
    columns,
):

    for col in range(
        1,
        columns + 1,
    ):

        cell = ws.cell(
            row=1,
            column=col,
        )

        cell.fill = HEAD
        cell.font = WHITE
        cell.border = BOX

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    ws.freeze_panes = "A2"


def style_body(ws):

    for row in ws.iter_rows(
        min_row=2,
    ):

        for cell in row:

            cell.font = BASE
            cell.border = BOX

            cell.alignment = Alignment(
                vertical="center",
            )


def color_status(cell):

    value = str(
        cell.value or ""
    ).upper()

    if (
        "SOBRECARGA" in value
        or
        "ESGOTADA" in value
    ):

        cell.fill = RED

    elif (
        "CRÍTICO" in value
        or
        "LIMITE" in value
        or
        "ALTA" in value
    ):

        cell.fill = YELLOW

    elif (
        value == "OK"
        or
        value == "NORMAL"
    ):

        cell.fill = GREEN


# ============================================================
# EXPORTAÇÃO
# ============================================================

def export(
    path,
    techs,
    assignments,
    stats,
    invalid,
    duplicates,
    unassigned,
    assignment_log,
):

    wb = Workbook()

    # ========================================================
    # PAINEL
    # ========================================================

    dashboard = wb.active

    dashboard.title = "Painel"

    dashboard["A1"] = (
        "CENTRAL DE VISTORIAS"
    )

    dashboard["A1"].font = Font(
        name="Arial",
        size=18,
        bold=True,
        color="FFFFFF",
    )

    dashboard["A1"].fill = HEAD

    dashboard.merge_cells(
        "A1:F2"
    )

    dashboard["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    indicators = [
        (
            "Total de sinistros",
            stats["total"],
        ),
        (
            "Distribuídos",
            stats["distributed"],
        ),
        (
            "Não distribuídos",
            stats["unassigned"],
        ),
        (
            "Taxa de distribuição",
            f'{stats["distribution_rate"]:.1f}%',
        ),
        (
            "Capacidade total",
            stats["capacity"],
        ),
    ]

    row = 4

    for label, value in indicators:

        dashboard.cell(
            row=row,
            column=1,
            value=label,
        )

        dashboard.cell(
            row=row,
            column=2,
            value=value,
        )

        dashboard.cell(
            row=row,
            column=1,
        ).font = BOLD

        dashboard.cell(
            row=row,
            column=1,
        ).border = BOX

        dashboard.cell(
            row=row,
            column=2,
        ).border = BOX

        row += 1

    dashboard["D4"] = (
        "Situação"
    )

    dashboard["D4"].font = BOLD

    if stats["unassigned"]:

        dashboard["E4"] = (
            "ATENÇÃO: existem "
            "sinistros não distribuídos"
        )

        dashboard["E4"].fill = RED

    elif invalid or duplicates:

        dashboard["E4"] = (
            "ATENÇÃO: existem "
            "linhas com problemas"
        )

        dashboard["E4"].fill = YELLOW

    else:

        dashboard["E4"] = (
            "DISTRIBUIÇÃO CONCLUÍDA"
        )

        dashboard["E4"].fill = GREEN

    dashboard.column_dimensions[
        "A"
    ].width = 28

    dashboard.column_dimensions[
        "B"
    ].width = 20

    dashboard.column_dimensions[
        "D"
    ].width = 18

    dashboard.column_dimensions[
        "E"
    ].width = 45

    # ========================================================
    # AGENDA
    # ========================================================

    agenda = wb.create_sheet(
        "Agenda"
    )

    agenda.append(
        [
            "HORÁRIO",
            "TÉCNICO",
            "TIPO",
            "SINISTRO",
            "ASSISTÊNCIA_COD",
        ]
    )

    style_header(
        agenda,
        5,
    )

    row_number = 2

    for time in TIMES:

        for tech in techs:

            name = tech[
                "name"
            ]

            for item in assignments[
                name
            ][time]:

                agenda.append(
                    [
                        time,
                        name,
                        tech["type"],
                        item["id"],
                        item["assist"],
                    ]
                )

                fill = (
                    ALT
                    if row_number % 2 == 0
                    else None
                )

                for col in range(
                    1,
                    6,
                ):

                    cell = agenda.cell(
                        row=row_number,
                        column=col,
                    )

                    cell.border = BOX

                    if fill:
                        cell.fill = fill

                row_number += 1

    agenda.column_dimensions[
        "A"
    ].width = 11

    agenda.column_dimensions[
        "B"
    ].width = 22

    agenda.column_dimensions[
        "C"
    ].width = 14

    agenda.column_dimensions[
        "D"
    ].width = 18

    agenda.column_dimensions[
        "E"
    ].width = 20

    # ========================================================
    # GRADE
    # ========================================================

    grade = wb.create_sheet(
        "Grade"
    )

    grade.append(
        ["HORÁRIO"]
        +
        [
            tech["name"]
            for tech in techs
        ]
    )

    style_header(
        grade,
        len(techs) + 1,
    )

    for row_number, time in enumerate(
        TIMES,
        start=2,
    ):

        grade.append(
            [
                time
            ]
            +
            [
                "\n".join(
                    item["id"]
                    for item in assignments[
                        tech["name"]
                    ][time]
                )
                for tech in techs
            ]
        )

        for col in range(
            1,
            len(techs) + 2,
        ):

            cell = grade.cell(
                row=row_number,
                column=col,
            )

            cell.border = BOX

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    grade.column_dimensions[
        "A"
    ].width = 11

    for col in range(
        2,
        len(techs) + 2,
    ):

        grade.column_dimensions[
            get_column_letter(col)
        ].width = 18

    # ========================================================
    # RESUMO
    # ========================================================

    summary = wb.create_sheet(
        "Resumo"
    )

    summary.append(
        [
            "TÉCNICO",
            "TIPO",
            "CAPACIDADE",
            "VISTORIAS",
            "RESTANTE",
            "OCUPAÇÃO",
            "STATUS",
        ]
    )

    style_header(
        summary,
        7,
    )

    for tech in techs:

        name = tech[
            "name"
        ]

        info = stats[
            "technicians"
        ][name]

        utilization = info[
            "utilization"
        ]

        if utilization >= 100:

            status = (
                "CAPACIDADE ESGOTADA"
            )

        elif utilization >= 90:

            status = (
                "ALTA OCUPAÇÃO"
            )

        elif utilization >= 75:

            status = (
                "OCUPAÇÃO MODERADA"
            )

        else:

            status = "NORMAL"

        summary.append(
            [
                name,
                tech["type"],
                info["capacity"],
                info["used"],
                info["remaining"],
                f"{utilization:.1f}%",
                status,
            ]
        )

    total_row = (
        len(techs) + 2
    )

    summary.append(
        [
            "TOTAL",
            "",
            f"=SUM(C2:C{total_row - 1})",
            f"=SUM(D2:D{total_row - 1})",
            f"=SUM(E2:E{total_row - 1})",
            "",
            "",
        ]
    )

    style_body(
        summary
    )

    for cell in summary[
        total_row
    ]:

        cell.font = BOLD
        cell.border = BOX

    for row in summary.iter_rows(
        min_row=2,
        max_col=7,
    ):

        color_status(
            row[6]
        )

    for column, width in {
        "A": 22,
        "B": 14,
        "C": 14,
        "D": 13,
        "E": 12,
        "F": 12,
        "G": 25,
    }.items():

        summary.column_dimensions[
            column
        ].width = width

    # ========================================================
    # CAPACIDADE
    # ========================================================

    capacity_sheet = wb.create_sheet(
        "Capacidade"
    )

    capacity_sheet.append(
        [
            "HORÁRIO",
            "DEMANDA",
            "TÉCNICOS",
            "CAPACIDADE",
            "SALDO",
            "OCUPAÇÃO",
            "STATUS",
        ]
    )

    style_header(
        capacity_sheet,
        7,
    )

    for time in TIMES:

        info = stats[
            "times"
        ][time]

        capacity_sheet.append(
            [
                time,
                info["demand"],
                info["available"],
                info["capacity"],
                info["remaining"],
                f'{info["utilization"]:.1f}%',
                info["status"],
            ]
        )

    style_body(
        capacity_sheet
    )

    for row in capacity_sheet.iter_rows(
        min_row=2,
        max_col=7,
    ):

        color_status(
            row[6]
        )

    for column, width in {
        "A": 12,
        "B": 12,
        "C": 14,
        "D": 14,
        "E": 12,
        "F": 12,
        "G": 18,
    }.items():

        capacity_sheet.column_dimensions[
            column
        ].width = width

    # ========================================================
    # NÃO DISTRIBUÍDAS
    # ========================================================

    if unassigned:

        sheet = wb.create_sheet(
            "Não distribuídas"
        )

        sheet.append(
            [
                "LINHA",
                "SINISTRO",
                "HORÁRIO",
                "ASSISTÊNCIA_COD",
                "MOTIVO",
            ]
        )

        style_header(
            sheet,
            5,
        )

        for item in unassigned:

            sheet.append(
                [
                    item["line"],
                    item["id"],
                    item["time"],
                    item["assist"],
                    item["reason"],
                ]
            )

        style_body(
            sheet
        )

        for column, width in {
            "A": 10,
            "B": 18,
            "C": 12,
            "D": 20,
            "E": 40,
        }.items():

            sheet.column_dimensions[
                column
            ].width = width

    # ========================================================
    # DUPLICADOS
    # ========================================================

    if duplicates:

        sheet = wb.create_sheet(
            "Duplicados"
        )

        sheet.append(
            [
                "LINHA",
                "SINISTRO",
                "HORÁRIO",
                "ASSISTÊNCIA_COD",
                "MOTIVO",
            ]
        )

        style_header(
            sheet,
            5,
        )

        for item in duplicates:

            sheet.append(
                [
                    item["line"],
                    item["id"],
                    item["time"],
                    item["assist"],
                    item["reason"],
                ]
            )

        style_body(
            sheet
        )

        for column, width in {
            "A": 10,
            "B": 18,
            "C": 12,
            "D": 20,
            "E": 30,
        }.items():

            sheet.column_dimensions[
                column
            ].width = width

    # ========================================================
    # LINHAS INVÁLIDAS
    # ========================================================

    if invalid:

        sheet = wb.create_sheet(
            "Linhas inválidas"
        )

        sheet.append(
            [
                "LINHA",
                "SINISTRO",
                "HORÁRIO",
                "ASSISTÊNCIA_COD",
                "MOTIVO",
            ]
        )

        style_header(
            sheet,
            5,
        )

        for item in invalid:

            sheet.append(
                [
                    item["line"],
                    item["id"],
                    item["time"],
                    item["assist"],
                    item["reason"],
                ]
            )

        style_body(
            sheet
        )

        for column, width in {
            "A": 10,
            "B": 18,
            "C": 12,
            "D": 20,
            "E": 40,
        }.items():

            sheet.column_dimensions[
                column
            ].width = width

    # ========================================================
    # AUDITORIA
    # ========================================================

    audit = wb.create_sheet(
        "Auditoria"
    )

    audit.append(
        [
            "LINHA",
            "SINISTRO",
            "HORÁRIO",
            "STATUS",
            "TÉCNICO",
            "MOTIVO",
            "SCORE",
        ]
    )

    style_header(
        audit,
        7,
    )

    for entry in assignment_log:

        item = entry[
            "item"
        ]

        audit.append(
            [
                item["line"],
                item["id"],
                item["time"],
                entry["status"],
                entry["tech"],
                entry["reason"],
                (
                    round(
                        entry["score"],
                        3,
                    )
                    if "score" in entry
                    else ""
                ),
            ]
        )

    style_body(
        audit
    )

    for row in audit.iter_rows(
        min_row=2,
        max_col=7,
    ):

        status_cell = row[3]

        if (
            status_cell.value
            == "NÃO DISTRIBUÍDO"
        ):

            status_cell.fill = RED

        else:

            status_cell.fill = GREEN

    for column, width in {
        "A": 10,
        "B": 18,
        "C": 12,
        "D": 20,
        "E": 22,
        "F": 35,
        "G": 12,
    }.items():

        audit.column_dimensions[
            column
        ].width = width

    # ========================================================
    # SALVAR
    # ========================================================

    wb.save(
        path
    )


# ============================================================
# MAIN
# ============================================================

def main():

    if len(sys.argv) > 3:

        print(
            "Uso:"
        )

        print(
            "python distribuir.py "
            "entrada.xlsx saida.xlsx"
        )

        return 1

    entrada = (
        sys.argv[1]
        if len(sys.argv) >= 2
        else "Entrada.xlsx"
    )

    saida = (
        sys.argv[2]
        if len(sys.argv) >= 3
        else "saida.xlsx"
    )

    # --------------------------------------------------------
    # Verifica arquivo de entrada.
    # --------------------------------------------------------

    if not Path(
        entrada
    ).exists():

        print()

        print(
            f"ERRO: arquivo não encontrado:"
            f" {entrada}"
        )

        return 1

    # --------------------------------------------------------
    # Valida técnicos.
    # --------------------------------------------------------

    tech_errors = validate_techs(
        TECHS
    )

    if tech_errors:

        print()

        print(
            "ERROS NA CONFIGURAÇÃO:"
        )

        for error in tech_errors:

            print(
                f"  - {error}"
            )

        return 1

    techs = get_active_techs()

    if not techs:

        print(
            "ERRO: nenhum técnico ativo."
        )

        return 1

    # --------------------------------------------------------
    # Processamento.
    # --------------------------------------------------------

    try:

        (
            items,
            invalid,
            duplicates,
        ) = parse_input(
            entrada
        )

        if not items:

            print()

            print(
                "ERRO: nenhum sinistro válido "
                "foi encontrado."
            )

            return 1

        (
            assignments,
            load,
            capacity,
            demand,
            available_techs,
            unassigned,
            assignment_log,
        ) = distribute(
            items,
            techs,
        )

        stats = calculate_stats(
            techs,
            load,
            capacity,
            demand,
            available_techs,
            unassigned,
        )

        export(
            saida,
            techs,
            assignments,
            stats,
            invalid,
            duplicates,
            unassigned,
            assignment_log,
        )

    except PermissionError:

        print()

        print(
            "ERRO: não foi possível salvar "
            "a planilha."
        )

        print(
            "Feche o arquivo de saída no Excel "
            "e tente novamente."
        )

        return 1

    except Exception as error:

        print()

        print(
            "ERRO durante o processamento:"
        )

        print(
            f"  {error}"
        )

        return 1

    # ========================================================
    # RESULTADO NO TERMINAL
    # ========================================================

    print()
    print("=" * 70)
    print(" CENTRAL DE VISTORIAS")
    print(" DISTRIBUIÇÃO AUTOMÁTICA")
    print("=" * 70)
    print()

    print(
        f"Sinistros válidos:       "
        f"{stats['total']}"
    )

    print(
        f"Distribuídos:             "
        f"{stats['distributed']}"
    )

    print(
        f"Não distribuídos:         "
        f"{stats['unassigned']}"
    )

    print(
        f"Duplicados:               "
        f"{len(duplicates)}"
    )

    print(
        f"Linhas inválidas:         "
        f"{len(invalid)}"
    )

    print(
        f"Capacidade total:         "
        f"{stats['capacity']}"
    )

    print(
        f"Taxa de distribuição:     "
        f"{stats['distribution_rate']:.1f}%"
    )

    print()

    print(
        "CARGA DOS TÉCNICOS"
    )

    print(
        "-" * 70
    )

    for tech in techs:

        name = tech[
            "name"
        ]

        info = stats[
            "technicians"
        ][name]

        print(
            f"{name:<22}"
            f" {info['used']:>3}"
            f"/{info['capacity']:<3}"
            f"  {info['utilization']:>6.1f}%"
        )

    # --------------------------------------------------------
    # Horários críticos.
    # --------------------------------------------------------

    critical_times = [
        time
        for time, info
        in stats[
            "times"
        ].items()
        if info["status"]
        in {
            "CRÍTICO",
            "NO LIMITE",
            "SOBRECARGA",
        }
    ]

    if critical_times:

        print()

        print(
            "HORÁRIOS CRÍTICOS:"
        )

        for time in critical_times:

            info = stats[
                "times"
            ][time]

            print(
                f"  {time}: "
                f"{info['demand']} demanda / "
                f"{info['capacity']} capacidade "
                f"-> {info['status']}"
            )

    # --------------------------------------------------------
    # Não distribuídos.
    # --------------------------------------------------------

    if unassigned:

        print()

        print(
            "NÃO DISTRIBUÍDOS:"
        )

        for item in unassigned:

            print(
                f"  Linha {item['line']}: "
                f"{item['id']} "
                f"às {item['time']} "
                f"- {item['reason']}"
            )

    # --------------------------------------------------------
    # Duplicados.
    # --------------------------------------------------------

    if duplicates:

        print()

        print(
            "DUPLICADOS:"
        )

        for item in duplicates:

            print(
                f"  Linha {item['line']}: "
                f"{item['id']}"
            )

    print()

    print(
        f"Planilha gerada: {saida}"
    )

    print()

    return 0


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":

    sys.exit(
        main()
    )